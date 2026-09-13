# Copyright (C) 2026 Matteo Sassi
# SPDX-License-Identifier: AGPL-3.0-only

"""Portable game-library files, independent of the graphical interface."""

import csv
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook, load_workbook

from . import copy_identifiers
from . import database as data


LEGACY_CSV_COLUMNS = ("game_name", "owner_label", "quantity")
CSV_COLUMNS = LEGACY_CSV_COLUMNS + ("copy_identifier", "identifier_source")


class ImportNotValid(Exception):
    pass


class LibraryResetBlocked(Exception):
    pass


@dataclass(frozen=True)
class ImportProblem:
    row_number: int
    message: str


@dataclass(frozen=True)
class ImportPreview:
    event_id: int
    rows: tuple[dict, ...]
    problems: tuple[ImportProblem, ...]
    existing_titles: tuple[str, ...]
    new_owner_labels: tuple[str, ...]

    @property
    def valid_count(self):
        return len(self.rows)

    @property
    def invalid_count(self):
        return len(self.problems)

    @property
    def can_apply(self):
        return bool(self.rows) and not self.problems


@dataclass(frozen=True)
class ImportResult:
    rows: int
    games_created: int
    owner_labels_created: int
    copies_created: int


def legacy_library_rows():
    """Read the aggregate legacy library without changing its database."""
    return [dict(row) for row in data.righe_export_ludoteca_legacy()]


def write_library_csv(path, rows, *, columns=CSV_COLUMNS):
    """Write canonical, comma-delimited UTF-8 library data."""
    with Path(path).open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns)
        writer.writeheader()
        writer.writerows(rows)


def write_library_xlsx(path, rows, *, columns=CSV_COLUMNS):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "game_library"
    sheet.append(columns)
    for row in rows:
        sheet.append(tuple(row.get(column, "") for column in columns))
    workbook.save(path)


def export_legacy_library_csv(path):
    rows = legacy_library_rows()
    write_library_csv(path, rows, columns=LEGACY_CSV_COLUMNS)
    return len(rows)


def event_library_rows(event_id):
    return [dict(row) for row in data.righe_export_ludoteca_evento(event_id)]


def export_event_library(path, event_id):
    rows = event_library_rows(event_id)
    destination = Path(path)
    if destination.suffix.casefold() == ".xlsx":
        write_library_xlsx(destination, rows)
    else:
        write_library_csv(destination, rows)
    return len(rows)


def _read_csv(path):
    with Path(path).open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        return tuple(reader.fieldnames or ()), list(reader)


def _read_xlsx(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        values = sheet.iter_rows(values_only=True)
        header = tuple(str(value).strip() if value is not None else ""
                       for value in next(values, ()))
        rows = [dict(zip(header, values_row)) for values_row in values]
        return header, rows
    finally:
        workbook.close()


def _read_rows(path):
    if Path(path).suffix.casefold() == ".xlsx":
        return _read_xlsx(path)
    return _read_csv(path)


def import_requires_owner(path):
    header, _ = _read_rows(path)
    return "owner_label" not in set(header)


def _normalized(value):
    return str(value).strip().casefold()


def preview_import(path, event_id, *, owner_label=None):
    header, source_rows = _read_rows(path)
    header_set = set(header)
    problems = []
    if not {"game_name", "quantity"}.issubset(header_set):
        problems.append(ImportProblem(1, "Missing game_name or quantity column"))
    has_owner_column = "owner_label" in header_set
    chosen_owner = str(owner_label).strip() if owner_label is not None else ""
    if not has_owner_column and not chosen_owner:
        problems.append(ImportProblem(1, "An owner label is required"))

    existing_games = {
        row["nome"].casefold(): row["nome"]
        for row in data.elenco_giochi_evento_backoffice(event_id)
    }
    existing_owners = {
        row["nome"].casefold(): row["nome"]
        for row in data.elenco_owner_labels(event_id)
    }
    valid_rows = []
    existing_titles = set()
    new_owners = set()
    known_identifiers = data.valori_identificatori_copie_evento(event_id)
    file_identifiers = set()
    if not problems:
        for row_number, source in enumerate(source_rows, start=2):
            game_name = str(source.get("game_name") or "").strip()
            owner_name = (
                str(source.get("owner_label") or "").strip()
                if has_owner_column else chosen_owner
            )
            try:
                quantity = int(source.get("quantity"))
                if quantity <= 0:
                    raise ValueError
            except (TypeError, ValueError):
                problems.append(ImportProblem(row_number, "Quantity must be positive"))
                continue
            if not game_name:
                problems.append(ImportProblem(row_number, "Game name is required"))
                continue
            if not owner_name:
                problems.append(ImportProblem(row_number, "Owner label is required"))
                continue
            identifier = ""
            identifier_source = ""
            raw_identifier = source.get("copy_identifier", "")
            if raw_identifier is not None and str(raw_identifier).strip():
                try:
                    identifier = copy_identifiers.normalize_identifier(
                        raw_identifier
                    )
                except copy_identifiers.InvalidCopyIdentifier as exc:
                    problems.append(ImportProblem(row_number, str(exc)))
                    continue
                raw_source = source.get("identifier_source", "")
                identifier_source = str(raw_source or "").strip().casefold()
                if not identifier_source:
                    identifier_source = "external"
                if identifier_source not in {"external", "ludox"}:
                    problems.append(ImportProblem(
                        row_number,
                        "Identifier source must be external or ludox",
                    ))
                    continue
                if quantity != 1:
                    problems.append(ImportProblem(
                        row_number,
                        "An identified copy must have quantity 1",
                    ))
                    continue
                if identifier in file_identifiers:
                    problems.append(ImportProblem(
                        row_number,
                        "Copy identifier is duplicated in the import file",
                    ))
                    continue
                if identifier in known_identifiers:
                    problems.append(ImportProblem(
                        row_number,
                        "Copy identifier already exists in this Event",
                    ))
                    continue
                file_identifiers.add(identifier)
            elif str(source.get("identifier_source") or "").strip():
                problems.append(ImportProblem(
                    row_number,
                    "Identifier source requires a copy identifier",
                ))
                continue
            game_key = _normalized(game_name)
            owner_key = _normalized(owner_name)
            if game_key in existing_games:
                existing_titles.add(existing_games[game_key])
                game_name = existing_games[game_key]
            if owner_key in existing_owners:
                owner_name = existing_owners[owner_key]
            else:
                new_owners.add(owner_name)
            valid_rows.append({
                "game_name": game_name,
                "game_key": game_key,
                "owner_label": owner_name,
                "owner_key": owner_key,
                "quantity": quantity,
                "copy_identifier": identifier,
                "identifier_source": identifier_source,
            })
    return ImportPreview(
        event_id,
        tuple(valid_rows),
        tuple(problems),
        tuple(sorted(existing_titles, key=str.casefold)),
        tuple(sorted(new_owners, key=str.casefold)),
    )


def apply_import(preview):
    if not preview.can_apply:
        raise ImportNotValid()
    games, owners, copies = data.applica_import_ludoteca(
        preview.event_id, preview.rows
    )
    return ImportResult(len(preview.rows), games, owners, copies)


def reset_event_library(event_id):
    if not data.reset_ludoteca_evento(event_id):
        raise LibraryResetBlocked()
